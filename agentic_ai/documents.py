"""Membaca isi file kantor (PDF, Word, Excel, CSV, teks) menjadi teks untuk agen."""

from __future__ import annotations

import csv
from pathlib import Path

TEXT_SUFFIXES = {".txt", ".md"}
DOCUMENT_SUFFIXES = {".pdf", ".docx"} | TEXT_SUFFIXES
SPREADSHEET_SUFFIXES = {".xlsx", ".xlsm", ".csv"}
SUPPORTED_SUFFIXES = DOCUMENT_SUFFIXES | SPREADSHEET_SUFFIXES


def document_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        pages = PdfReader(path).pages
        return "\n\n".join(f"[Halaman {i}]\n{p.extract_text() or ''}" for i, p in enumerate(pages, 1))
    if suffix == ".docx":
        import docx

        doc = docx.Document(path)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            parts.extend(" | ".join(c.text.strip() for c in row.cells) for row in table.rows)
        return "\n".join(parts)
    if suffix in TEXT_SUFFIXES:
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix in SPREADSHEET_SUFFIXES:
        return "\n".join(
            f"[Sheet {name}]\n" + "\n".join(" | ".join(row) for row in rows)
            for name, rows in sheets(path).items()
        )
    raise ValueError(f"Format {suffix} belum didukung. Didukung: {', '.join(sorted(SUPPORTED_SUFFIXES))}")


def sheets(path: Path) -> dict[str, list[list[str]]]:
    """Semua sheet sebagai daftar baris (baris pertama biasanya header)."""
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as f:
            return {path.stem: [row for row in csv.reader(f)]}
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return {
            ws.title: [[_cell(v) for v in row] for row in ws.iter_rows(values_only=True)
                       if any(v is not None for v in row)]
            for ws in wb.worksheets
        }
    finally:
        wb.close()


def _cell(value: object) -> str:
    if value is None:
        return ""
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d")
    return str(value)
