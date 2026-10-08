"""Membuat data HRD contoh (fiktif) di contoh_data/HRD untuk mencoba agen tanpa data asli.

Jalankan: python scripts/buat_contoh_data.py
"""

from __future__ import annotations

import datetime as dt
import random
from pathlib import Path

import docx
import openpyxl

OUT = Path(__file__).resolve().parent.parent / "contoh_data" / "HRD"

KARYAWAN = [
    # NIK, nama, bagian, jabatan, shift, tanggal masuk, status
    ("K001", "Budi Santoso", "Produksi - Extruder", "Operator", "Shift", "2019-03-11", "Tetap"),
    ("K002", "Siti Aminah", "Produksi - Printing", "Operator", "Shift", "2021-07-01", "Tetap"),
    ("K003", "Agus Prasetyo", "Produksi - Cutting", "Kepala Regu", "Shift", "2017-01-16", "Tetap"),
    ("K004", "Dewi Lestari", "Quality Control", "QC Inspector", "Shift", "2022-02-14", "Kontrak"),
    ("K005", "Rudi Hartono", "Gudang", "Admin Gudang", "Non-shift", "2020-09-07", "Tetap"),
    ("K006", "Rina Wulandari", "HRD", "Staf HRD", "Non-shift", "2023-05-02", "Kontrak"),
    ("K007", "Joko Susilo", "Produksi - Extruder", "Teknisi Mesin", "Shift", "2018-11-19", "Tetap"),
    ("K008", "Maya Sari", "Sales", "Sales Executive", "Non-shift", "2024-01-08", "Kontrak"),
]

SHIFTS = {"Shift 1": ("07:00", "15:00"), "Shift 2": ("15:00", "23:00"), "Shift 3": ("23:00", "07:00"),
          "Non-shift": ("08:00", "17:00")}


def karyawan_xlsx() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Karyawan"
    ws.append(["NIK", "Nama", "Bagian", "Jabatan", "Pola Kerja", "Tanggal Masuk", "Status"])
    for row in KARYAWAN:
        ws.append([*row[:5], dt.date.fromisoformat(row[5]), row[6]])
    wb.save(OUT / "Data_Karyawan.xlsx")


def absensi_xlsx() -> None:
    rng = random.Random(7)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "September 2026"
    ws.append(["Tanggal", "NIK", "Nama", "Shift", "Jam Masuk", "Jam Pulang", "Keterangan", "Lembur (jam)"])
    day = dt.date(2026, 9, 1)
    while day.month == 9:
        if day.weekday() < 6:  # Senin-Sabtu
            for i, (nik, nama, _, _, pola, _, _) in enumerate(KARYAWAN):
                shift = f"Shift {(i + day.day // 7) % 3 + 1}" if pola == "Shift" else "Non-shift"
                masuk, pulang = SHIFTS[shift]
                roll = rng.random()
                if roll < 0.03:
                    ws.append([day, nik, nama, shift, "", "", "Sakit (surat dokter)", 0])
                elif roll < 0.05:
                    ws.append([day, nik, nama, shift, "", "", "Alpa", 0])
                elif roll < 0.08:
                    ws.append([day, nik, nama, shift, "", "", "Cuti tahunan", 0])
                else:
                    h, m = map(int, masuk.split(":"))
                    late = rng.choice([0] * 8 + [7, 18, 35])
                    jam_masuk = f"{h:02d}:{m + late:02d}" if late < 60 else masuk
                    lembur = rng.choice([0] * 6 + [1, 2, 3])
                    ws.append([day, nik, nama, shift, jam_masuk, pulang,
                               "Terlambat" if late else "Hadir", lembur])
        day += dt.timedelta(days=1)
    wb.save(OUT / "Absensi_2026-09.xlsx")


def cuti_xlsx() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cuti 2026"
    ws.append(["NIK", "Nama", "Hak Cuti Tahunan", "Cuti Terpakai", "Pengajuan Terakhir", "Status Pengajuan"])
    terpakai = [5, 8, 12, 3, 6, 2, 9, 1]
    for (nik, nama, *_), used in zip(KARYAWAN, terpakai):
        ws.append([nik, nama, 12, used, dt.date(2026, 9, 20), "Disetujui" if used < 12 else "Ditolak"])
    wb.save(OUT / "Rekap_Cuti_2026.xlsx")


def peraturan_docx() -> None:
    d = docx.Document()
    d.add_heading("PERATURAN PERUSAHAAN PT CONTOH PLASTIK (FIKTIF)", 0)
    pasal = [
        ("Pasal 5 - Jam Kerja",
         "Karyawan non-shift bekerja Senin-Jumat 08.00-17.00 dengan istirahat 1 jam, dan Sabtu 08.00-13.00. "
         "Karyawan produksi bekerja dalam 3 shift: Shift 1 07.00-15.00, Shift 2 15.00-23.00, Shift 3 23.00-07.00. "
         "Pergantian shift dilakukan setiap minggu."),
        ("Pasal 6 - Keterlambatan",
         "Karyawan dianggap terlambat bila hadir lebih dari 5 menit setelah jam mulai. "
         "Keterlambatan 3 kali dalam satu bulan mendapat teguran lisan dari atasan; "
         "5 kali atau lebih dalam satu bulan mendapat Surat Peringatan (SP) 1."),
        ("Pasal 7 - Ketidakhadiran",
         "Tidak masuk tanpa keterangan (alpa) 1 hari mendapat teguran tertulis. "
         "Alpa 3 hari dalam satu bulan mendapat SP 1. Sakit wajib disertai surat dokter."),
        ("Pasal 9 - Cuti Tahunan",
         "Karyawan yang telah bekerja 12 bulan berturut-turut berhak atas cuti tahunan 12 hari kerja. "
         "Pengajuan cuti paling lambat 7 hari sebelumnya melalui atasan langsung dan disetujui HRD. "
         "Untuk bagian produksi, cuti diatur agar tiap shift tetap memenuhi jumlah operator minimum. "
         "Sisa cuti tidak dapat diuangkan dan hangus pada akhir tahun berjalan."),
        ("Pasal 10 - Lembur",
         "Lembur harus atas perintah tertulis atasan. Lembur maksimal 4 jam per hari dan 18 jam per minggu."),
        ("Pasal 14 - Keselamatan Kerja (K3)",
         "Operator wajib memakai APD: sarung tangan tahan panas di area extruder, penutup telinga di area mesin, "
         "dan sepatu keselamatan di seluruh area produksi. Pelanggaran APD mendapat teguran tertulis."),
        ("Pasal 18 - Surat Peringatan",
         "SP berlaku 6 bulan. SP 1, SP 2, dan SP 3 diberikan berurutan untuk pelanggaran yang berulang "
         "dalam masa berlaku SP sebelumnya. SP ditandatangani manajer HRD."),
    ]
    for judul, isi in pasal:
        d.add_heading(judul, level=2)
        d.add_paragraph(isi)
    d.save(OUT / "Peraturan_Perusahaan.docx")


def sop_pdf() -> None:
    lines = [
        "SOP-HRD-03 Pengajuan Cuti (FIKTIF)",
        "1. Karyawan mengisi formulir cuti minimal 7 hari sebelum tanggal cuti.",
        "2. Atasan langsung memeriksa jadwal shift dan menyetujui atau menolak dalam 2 hari kerja.",
        "3. HRD memeriksa sisa hak cuti di Rekap Cuti, lalu mencatat persetujuan akhir.",
        "4. Cuti mendadak karena keadaan darurat dilaporkan ke atasan paling lambat hari itu juga.",
    ]
    _write_simple_pdf(OUT / "SOP_Pengajuan_Cuti.pdf", lines)


def _write_simple_pdf(path: Path, lines: list[str]) -> None:
    """PDF satu halaman berisi teks polos, tanpa dependensi tambahan."""
    text = "BT /F1 11 Tf 50 800 Td 16 TL " + " ".join(
        "(" + ln.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)") + ") Tj T*" for ln in lines
    ) + " ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R "
        "/Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(text)} >>\nstream\n{text}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = "%PDF-1.4\n"
    offsets = []
    for i, obj in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{obj}\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n"
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    path.write_bytes(out.encode("latin-1"))


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    karyawan_xlsx()
    absensi_xlsx()
    cuti_xlsx()
    peraturan_docx()
    sop_pdf()
    print(f"Data contoh dibuat di {OUT}")
